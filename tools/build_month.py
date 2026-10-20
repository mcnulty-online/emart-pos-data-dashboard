#!/usr/bin/env python3
"""
이마트 POS 원본 → 대시보드 데이터 파일 변환기

  python tools/build_month.py 원본파일.xlsx [원본파일2.csv ...]

하는 일
  1) 원본(엑셀/CSV)에서 일자·점포·상품·수량·금액 열을 찾아 (점포, 상품, 일자)별로 합산
  2) data/YYYY-MM.js 를 만들거나 갱신 (원본에 있는 날짜만 교체, 나머지 날짜는 유지)
  3) 새 점포·상품을 data/master.js 에 추가 (상품 분류가 원본에 없으면 '미분류'로 넣고 경고)
  4) data/manifest.js 의 월 목록과 갱신일을 업데이트

원본 형식: 한 행 = 하루·점포·상품 (세로형). 열 이름은 아래 ALIASES 중 하나면 자동 인식합니다.
          제목 줄이 위에 몇 줄 있어도 머리글 행을 자동으로 찾습니다.
옵션: --data 폴더 경로(기본: 이 파일 기준 ../data)  --replace-month  원본에 없는 날짜도 지우고 그 달을 통째로 교체
"""
import argparse, csv, datetime as dt, json, re, sys
from pathlib import Path

ALIASES = {
    'date':  ['일자', '판매일자', '매출일자', '영업일자', '날짜', '일별', 'date'],
    'scode': ['점포코드', '점코드', '점포번호', '점번호'],
    'sname': ['점포명', '점명', '점포'],
    'pcode': ['상품코드', '바코드', '상품바코드', '단품코드', 'barcode'],
    'pname': ['상품명', '단품명', '품명'],
    'qty':   ['판매수량', '매출수량', '수량'],
    'amt':   ['판매금액', '매출금액', '매입금액', '금액'],
    'mid':   ['중분류', '중분류명'],
    'sub':   ['소분류', '소분류명'],
    'brand': ['브랜드', '브랜드명'],
}
REQUIRED = ['date', 'scode', 'pcode', 'qty', 'amt']
norm = lambda s: re.sub(r'[\s_()\[\]]', '', str(s or '')).lower()

# ---------------- 데이터 파일 읽기/쓰기 (JSONP 형식: EMART_XXX({...});) ----------------
def read_js(path):
    t = Path(path).read_text(encoding='utf-8')
    i, j = t.index('(', t.index('EMART_')), t.rindex(')')
    return json.loads(t[i + 1:j])

def dumps(o): return json.dumps(o, ensure_ascii=False, separators=(',', ':'))

def write_month(path, m):
    rows = sorted(m['rows'], key=lambda r: (str(r[0]), str(r[1])))
    body = ',\n'.join(dumps(r) for r in rows)
    Path(path).write_text(
        f"// {m['month']} 이마트 POS 일별 판매. rows: [점포코드, 상품코드, [일(1~31), 수량, 금액, ...]]  금액 = 이마트 매입원가 기준 판매금액\n"
        f"EMART_MONTH({{\"month\":\"{m['month']}\",\"days\":{m['days']},\"rows\":[\n{body}\n]}});\n", encoding='utf-8')

def write_master(path, ms):
    st = ',\n'.join(dumps(s) for s in ms['stores']); pr = ',\n'.join(dumps(p) for p in ms['prods'])
    Path(path).write_text(
        "// 점포·상품 마스터 + 공휴일. tools/build_month.py 가 새 점포/상품을 자동 추가합니다.\n"
        "// stores: [점포코드, 점포명, 점포유형]  prods: [상품코드, 상품명, 중분류, 소분류, 브랜드]  holidays: {'YYYY-MM-DD': '이름'}\n"
        "// '미분류' 상품은 중분류·소분류·브랜드를 직접 채워 주세요.\n"
        f"EMART_MASTER({{\n\"stores\":[\n{st}\n],\n\"prods\":[\n{pr}\n],\n\"holidays\":{dumps(ms.get('holidays', {}))}\n}});\n", encoding='utf-8')

def write_manifest(path, months):
    Path(path).write_text(
        "// 불러올 월 목록. build_month.py 가 자동으로 관리합니다 (손으로 고칠 때는 'YYYY-MM' 순서대로).\n"
        f"EMART_MANIFEST({dumps({'months': sorted(set(months)), 'updated': dt.datetime.now().strftime('%Y-%m-%d %H:%M')})});\n",
        encoding='utf-8')

def store_type(name):
    if '가상점' in name: return '온라인 가상점'
    if name.startswith('SFM'): return '스타필드마켓(SFM)'
    if name.startswith('EMF'): return 'EMF·EMFC'
    if name.startswith('EMEO'): return '기타'
    if name.startswith('EM'): return '이마트(EM)'
    return '기타'

# ---------------- 원본 읽기 ----------------
def read_table(path):
    p = Path(path)
    if p.suffix.lower() in ('.xlsx', '.xlsm', '.xls'):
        try:
            import openpyxl
        except ImportError:
            sys.exit('엑셀을 읽으려면: pip install openpyxl')
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if find_header(rows)[0] is not None: return rows
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    for enc in ('utf-8-sig', 'cp949', 'euc-kr'):
        try:
            with open(p, encoding=enc, newline='') as f: return list(csv.reader(f))
        except UnicodeDecodeError:
            continue
    sys.exit(f'{p}: 인코딩을 알 수 없습니다')

def find_header(rows):
    for i, r in enumerate(rows[:30]):
        cols = {}
        names = [norm(c) for c in r]
        used = set()
        for exact in (True, False):            # 1차: 열 이름이 정확히 같은 것, 2차: 이름에 포함된 것 (이미 쓴 열은 제외)
            for key, al in ALIASES.items():
                if key in cols: continue
                for a in al:
                    hit = [j for j, n in enumerate(names) if n and j not in used and (n == norm(a) if exact else norm(a) in n)]
                    if hit: cols[key] = hit[0]; used.add(hit[0]); break
        if all(k in cols for k in REQUIRED): return i, cols
    return None, None

def to_date(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    s = re.sub(r'[^0-9]', '', str(v or ''))
    if len(s) >= 8: return dt.date(int(s[:4]), int(s[4:6]), int(s[6:8]))
    return None

def to_num(v):
    if v is None or v == '': return 0
    if isinstance(v, (int, float)): return v
    s = str(v).replace(',', '').strip()
    if s.startswith('(') and s.endswith(')'): s = '-' + s[1:-1]
    try: return float(s)
    except ValueError: return 0

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('files', nargs='+')
    ap.add_argument('--data', default=str(Path(__file__).resolve().parent.parent / 'data'))
    ap.add_argument('--replace-month', action='store_true')
    a = ap.parse_args()
    data = Path(a.data)
    master = read_js(data / 'master.js') if (data / 'master.js').exists() else {'stores': [], 'prods': [], 'holidays': {}}
    mf = data / 'manifest.js'
    months = read_js(mf)['months'] if mf.exists() else []
    S = {str(s[0]): s for s in master['stores']}; P = {str(p[0]): p for p in master['prods']}
    agg = {}                                  # (YYYY-MM) -> {(scode,pcode,day): [q,a]}
    new_s, new_p, unk = [], [], set()
    for fpath in a.files:
        rows = read_table(fpath)
        h, cols = find_header(rows)
        if h is None:
            sys.exit(f'{fpath}: 머리글을 찾지 못했습니다. 필요한 열: 일자·점포코드·상품코드·판매수량·판매금액 (현재 첫 줄: {rows[:1]})')
        print(f'{fpath}: 머리글 {h + 1}행 ·', ', '.join(f'{k}={rows[h][j]}' for k, j in cols.items()))
        n = 0
        for r in rows[h + 1:]:
            g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
            d = to_date(g('date')); sc = str(g('scode') or '').strip(); pc = str(g('pcode') or '').strip()
            if not d or not sc or not pc: continue
            sc = sc[:-2] if sc.endswith('.0') else sc; pc = pc[:-2] if pc.endswith('.0') else pc
            q, am = to_num(g('qty')), to_num(g('amt'))
            if q == 0 and am == 0: continue
            mk = f'{d.year}-{d.month:02d}'
            cell = agg.setdefault(mk, {}).setdefault((sc, pc, d.day), [0, 0])
            cell[0] += q; cell[1] += am; n += 1
            if sc not in S:
                nm = str(g('sname') or sc).strip(); S[sc] = [sc, nm, store_type(nm)]; master['stores'].append(S[sc]); new_s.append(nm)
            if pc not in P:
                mid, sub, br = (str(g(k) or '').strip() or '미분류' for k in ('mid', 'sub', 'brand'))
                P[pc] = [pc, str(g('pname') or pc).strip(), mid, sub, br]; master['prods'].append(P[pc]); new_p.append(P[pc][1])
                if '미분류' in (mid, sub, br): unk.add(P[pc][1])
        print(f'  유효 행 {n:,}개')
    if not agg: sys.exit('변환할 데이터가 없습니다')
    for mk, cells in sorted(agg.items()):
        y, m = map(int, mk.split('-')); days = ((dt.date(y + m // 12, m % 12 + 1, 1)) - dt.date(y, m, 1)).days
        path = data / f'{mk}.js'
        keep = {}
        new_days = {k[2] for k in cells}
        if path.exists() and not a.replace_month:
            for sc, pc, f in read_js(path)['rows']:
                for i in range(0, len(f), 3):
                    if f[i] not in new_days: keep[(str(sc), str(pc), f[i])] = [f[i + 1], f[i + 2]]
        keep.update(cells)
        rows = {}
        for (sc, pc, d), (q, am) in sorted(keep.items(), key=lambda x: x[0][2]):
            q = int(q) if float(q).is_integer() else q; am = int(round(am))
            rows.setdefault((sc, pc), []).extend([d, q, am])
        days = min(days, max(k[2] for k in keep))          # 진행 중인 달은 데이터가 있는 마지막 날까지만 (일평균·전월 동기간 비교용)
        write_month(path, {'month': mk, 'days': days, 'rows': [[sc, pc, f] for (sc, pc), f in rows.items()]})
        months.append(mk)
        dmin, dmax = min(new_days), max(new_days)
        print(f'→ {path.name}: {mk}-{dmin:02d} ~ {mk}-{dmax:02d} 반영 ({len(new_days)}일, 점포×상품 {len(rows):,}건)')
    write_master(data / 'master.js', master); write_manifest(mf, months)
    if new_s: print(f'새 점포 {len(new_s)}개 추가:', ', '.join(new_s))
    if new_p: print(f'새 상품 {len(new_p)}개 추가:', ', '.join(new_p))
    if unk: print('※ 분류가 비어 있는 상품 → data/master.js 에서 중분류·소분류·브랜드를 채워 주세요:', ', '.join(sorted(unk)))
    print('완료. data 폴더를 GitHub에 올리면 대시보드에 반영됩니다.')

if __name__ == '__main__':
    main()
